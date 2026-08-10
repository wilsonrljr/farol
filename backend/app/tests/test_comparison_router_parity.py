from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app.api.errors import PublicInputError
from backend.app.api.routers import simulations
from backend.app.main import app
from backend.app.models import ComparisonInput

client = TestClient(app)


def _comparison_payload(**overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "property_value": 100_000.0,
        "down_payment": 20_000.0,
        "total_savings": 120_000.0,
        "loan_term_years": 1,
        "annual_interest_rate": 9.6,
        "loan_type": "PRICE",
        "rent_value": 500.0,
        "investment_returns": [{"start_month": 1, "annual_rate": 0.0}],
        "monthly_plan": {
            "net_income": 20_000.0,
            "non_housing_expenses": 5_000.0,
            "adjust_for_inflation": False,
            "wealth_allocation_percentage": 80.0,
            "financed_purchase": {
                "amortization_percentage": 25.0,
                "amortization_effect": "reduce_term",
            },
        },
        "additional_costs": {
            "itbi_percentage": 0.0,
            "deed_percentage": 0.0,
            "monthly_hoa": 0.0,
            "monthly_property_tax": 0.0,
        },
        "inflation_rate": 0.0,
        "rent_inflation_rate": 0.0,
        "property_appreciation_rate": 0.0,
    }
    payload.update(overrides)
    return payload


def test_canonical_comparison_arguments_include_behavior_flags() -> None:
    input_data = ComparisonInput.model_validate(_comparison_payload())

    kwargs = simulations.build_comparison_kwargs(input_data)

    assert kwargs["monthly_plan"] == input_data.monthly_plan
    assert kwargs["extra_income_events"] == input_data.extra_income_events


def test_enhanced_response_preserves_basic_opportunity_cost() -> None:
    payload = _comparison_payload(
        investment_returns=[{"start_month": 1, "annual_rate": 12.0}],
    )

    basic_response = client.post("/api/compare-scenarios", json=payload)
    enhanced_response = client.post("/api/compare-scenarios-enhanced", json=payload)

    assert basic_response.status_code == 200, basic_response.text
    assert enhanced_response.status_code == 200, enhanced_response.text
    basic_buy = next(
        scenario
        for scenario in basic_response.json()["scenarios"]
        if scenario["scenario_type"] == "buy"
    )
    enhanced_buy = next(
        scenario
        for scenario in enhanced_response.json()["scenarios"]
        if scenario["scenario_type"] == "buy"
    )

    assert basic_buy["opportunity_cost"] > 0
    assert enhanced_buy["opportunity_cost"] == pytest.approx(
        basic_buy["opportunity_cost"]
    )


def test_opportunity_gain_excludes_monthly_investment_principal() -> None:
    input_data = ComparisonInput.model_validate(
        _comparison_payload(
            monthly_plan={
                "net_income": 20_000.0,
                "non_housing_expenses": 0.0,
                "adjust_for_inflation": False,
                "wealth_allocation_percentage": 100.0,
                "financed_purchase": {
                    "amortization_percentage": 0.0,
                    "amortization_effect": "reduce_term",
                },
            },
            investment_returns=[{"start_month": 1, "annual_rate": 0.0}],
        )
    )

    result = simulations.run_basic_comparison(input_data)
    buy = next(
        scenario for scenario in result.scenarios if scenario.scenario_type == "buy"
    )

    assert buy.total_investment_from_income > 0.0
    assert buy.opportunity_cost == pytest.approx(0.0)


def test_enhanced_xlsx_matches_api_and_long_shape_omits_wide_sheet() -> None:
    payload = _comparison_payload()
    api_response = client.post("/api/compare-scenarios-enhanced", json=payload)
    export_response = client.post(
        "/api/compare-scenarios-enhanced/export?format=xlsx&shape=long",
        json=payload,
    )

    assert api_response.status_code == 200, api_response.text
    assert export_response.status_code == 200, export_response.text

    api_scenario = next(
        scenario
        for scenario in api_response.json()["scenarios"]
        if scenario["scenario_type"] == "invest_buy"
    )
    api_contributions = sum(
        float(month.get("investment_allocation") or 0.0)
        for month in api_scenario["monthly_data"]
    )

    workbook = load_workbook(
        BytesIO(export_response.content), read_only=True, data_only=True
    )
    assert "monthly_long" in workbook.sheetnames
    assert "monthly_wide" not in workbook.sheetnames

    sheet = workbook["monthly_long"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    scenario_index = headers.index("scenario")
    contribution_index = headers.index("investment_allocation")
    export_contributions = sum(
        float(row[contribution_index] or 0.0)
        for row in rows
        if row[scenario_index] == api_scenario["name"]
    )

    assert export_contributions == api_contributions


def test_batch_is_atomic_instead_of_silently_dropping_failed_item(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    original = simulations.run_enhanced_comparison

    def _fail_second(input_data: ComparisonInput):  # type: ignore[no-untyped-def]
        if input_data.property_value == 200_000.0:
            raise PublicInputError("second preset failed")
        return original(input_data)

    monkeypatch.setattr(simulations, "run_enhanced_comparison", _fail_second)
    response = client.post(
        "/api/compare-scenarios-batch",
        json={
            "items": [
                {
                    "preset_id": "first",
                    "preset_name": "First",
                    "input": _comparison_payload(monthly_plan=None),
                },
                {
                    "preset_id": "second",
                    "preset_name": "Second",
                    "input": _comparison_payload(
                        property_value=200_000.0, monthly_plan=None
                    ),
                },
            ]
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"] == "second preset failed"


def test_batch_keeps_exploratory_results_without_inventing_global_winner() -> None:
    response = client.post(
        "/api/compare-scenarios-batch",
        json={
            "items": [
                {
                    "preset_id": "first",
                    "preset_name": "First",
                    "input": _comparison_payload(monthly_plan=None),
                },
                {
                    "preset_id": "second",
                    "preset_name": "Second",
                    "input": _comparison_payload(
                        property_value=200_000.0,
                        monthly_plan=None,
                    ),
                },
            ]
        },
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert len(data["results"]) == 2
    assert data.get("global_best") is None
    assert data["ranking"] == []
    assert data["comparison_status"] == "no_authoritative_result"
    assert data["warnings"]


def test_batch_omits_global_ranking_across_different_resource_baselines() -> None:
    common: dict[str, object] = {}
    response = client.post(
        "/api/compare-scenarios-batch",
        json={
            "items": [
                {
                    "preset_id": "first",
                    "preset_name": "First",
                    "input": _comparison_payload(**common),
                },
                {
                    "preset_id": "second",
                    "preset_name": "Second",
                    "input": _comparison_payload(
                        **common,
                        total_savings=150_000.0,
                    ),
                },
            ]
        },
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert all(
        item["result"]["comparison_status"] == "comparable" for item in data["results"]
    )
    assert data.get("global_best") is None
    assert data["ranking"] == []
    assert data["comparison_status"] == "no_authoritative_result"
    assert any("different initial resources" in warning for warning in data["warnings"])


def test_batch_ranks_every_feasible_scenario_with_a_common_resource_baseline() -> None:
    common: dict[str, object] = {}
    response = client.post(
        "/api/compare-scenarios-batch",
        json={
            "items": [
                {
                    "preset_id": "first",
                    "preset_name": "First",
                    "input": _comparison_payload(**common),
                },
                {
                    "preset_id": "second",
                    "preset_name": "Second",
                    "input": _comparison_payload(
                        **common,
                        property_value=200_000.0,
                    ),
                },
            ]
        },
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["comparison_status"] == "ranked"
    assert len(data["ranking"]) >= 5
    wealth = [entry["final_wealth"] for entry in data["ranking"]]
    assert wealth == sorted(wealth, reverse=True)
    assert data["global_best"] == data["ranking"][0]
    feasible_count = sum(
        scenario["is_feasible"] is True
        for item in data["results"]
        for scenario in item["result"]["scenarios"]
    )
    assert len(data["ranking"]) == feasible_count


def test_sensitivity_uses_honest_discrete_terms() -> None:
    assert simulations._sensitivity_values("loan_term_years", 1.0, 5.0, 3) == [
        1.0,
        3.0,
        5.0,
    ]

    with pytest.raises(ValueError, match="too narrow"):
        simulations._sensitivity_values("loan_term_years", 1.0, 2.0, 5)


def test_sensitivity_base_value_uses_effective_input_representations() -> None:
    monthly_rate_input = ComparisonInput.model_validate(
        _comparison_payload(
            annual_interest_rate=None,
            monthly_interest_rate=1.0,
        )
    )
    assert simulations._get_parameter_value(
        monthly_rate_input, "annual_interest_rate"
    ) == pytest.approx((1.01**12 - 1) * 100)

    percentage_rent_input = ComparisonInput.model_validate(
        _comparison_payload(
            rent_value=None,
            rent_percentage=0.5,
        )
    )
    assert simulations._get_parameter_value(
        percentage_rent_input, "rent_value"
    ) == pytest.approx(500.0)

    fallback_appreciation_input = ComparisonInput.model_validate(
        _comparison_payload(
            inflation_rate=6.0,
            property_appreciation_rate=None,
        )
    )
    assert simulations._get_parameter_value(
        fallback_appreciation_input, "property_appreciation_rate"
    ) == pytest.approx(6.0)


def test_sensitivity_rejects_partial_range_instead_of_omitting_points() -> None:
    response = client.post(
        "/api/sensitivity-analysis",
        json={
            "base_input": _comparison_payload(
                property_value=100_000.0,
                down_payment=0.0,
                total_savings=100_000.0,
            ),
            "parameter": "down_payment",
            "range": {"min_value": 0.0, "max_value": 150_000.0, "steps": 4},
        },
    )

    assert response.status_code == 400
    assert "down_payment" in response.json()["detail"]


def test_sensitivity_returns_all_exploratory_points_without_a_fake_winner() -> None:
    response = client.post(
        "/api/sensitivity-analysis",
        json={
            "base_input": _comparison_payload(monthly_plan=None),
            "parameter": "inflation_rate",
            "range": {"min_value": 0.0, "max_value": 2.0, "steps": 3},
        },
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert [point["parameter_value"] for point in data["data_points"]] == [
        0.0,
        1.0,
        2.0,
    ]
    assert all(point.get("best_scenario") is None for point in data["data_points"])
    assert all(len(point["scenarios"]) == 3 for point in data["data_points"])
    assert data.get("best_overall") is None
    assert data["comparison_status"] == "no_authoritative_result"
    assert data["warnings"]


def test_sensitivity_ranked_points_match_direct_comparisons() -> None:
    payload = _comparison_payload()
    response = client.post(
        "/api/sensitivity-analysis",
        json={
            "base_input": payload,
            "parameter": "annual_interest_rate",
            "range": {"min_value": 0.0, "max_value": 12.0, "steps": 3},
        },
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["comparison_status"] == "ranked"
    assert [point["parameter_value"] for point in data["data_points"]] == [
        0.0,
        6.0,
        12.0,
    ]
    assert data["best_overall"] is not None

    base_input = ComparisonInput.model_validate(payload)
    for point in data["data_points"]:
        modified = simulations._apply_parameter_value(
            base_input,
            "annual_interest_rate",
            point["parameter_value"],
        )
        direct = simulations.run_enhanced_comparison(modified)
        assert point["comparison_status"] == direct.comparison_status
        assert point["best_scenario"] == direct.best_scenario
        assert point["best_scenario_type"] == direct.best_scenario_type
        assert set(point["scenarios"]) == {
            scenario.name for scenario in direct.scenarios
        }
        for scenario in direct.scenarios:
            assert point["scenarios"][scenario.name]["final_wealth"] == pytest.approx(
                scenario.final_wealth
            )


def test_loan_term_sensitivity_does_not_rank_different_horizons() -> None:
    response = client.post(
        "/api/sensitivity-analysis",
        json={
            "base_input": _comparison_payload(),
            "parameter": "loan_term_years",
            "range": {"min_value": 1.0, "max_value": 3.0, "steps": 3},
        },
    )

    assert response.status_code == 200, response.text
    data = response.json()
    assert all(
        point["comparison_status"] == "comparable" for point in data["data_points"]
    )
    assert data.get("best_overall") is None
    assert data["breakeven_points"] == []
    assert data["comparison_status"] == "no_authoritative_result"
    assert any(
        "changes the simulation horizon" in warning for warning in data["warnings"]
    )


def test_return_sensitivity_rejects_an_ambiguous_multi_period_curve() -> None:
    response = client.post(
        "/api/sensitivity-analysis",
        json={
            "base_input": _comparison_payload(
                investment_returns=[
                    {"start_month": 1, "end_month": 6, "annual_rate": 4.0},
                    {"start_month": 7, "annual_rate": 8.0},
                ],
            ),
            "parameter": "investment_return_rate",
            "range": {"min_value": 3.0, "max_value": 9.0, "steps": 3},
        },
    )

    assert response.status_code == 400
    assert "exactly one return range" in response.json()["detail"]
